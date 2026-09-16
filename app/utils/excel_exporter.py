from dataclasses import fields
from enum import Enum
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font


class ExcelExporter:
    """Exports dataclass collections into a formatted XLSX workbook."""
    @staticmethod
    def export(collections:dict[str,list],path:str="exports/donor_events.xlsx")->Path:
        target=Path(path);target.parent.mkdir(parents=True,exist_ok=True)
        wb=Workbook();wb.remove(wb.active)
        for title,items in collections.items():
            ws=wb.create_sheet(title)
            if not items:ws.append(["Нет данных"]);continue
            headers=[f.name for f in fields(items[0])];ws.append(headers)
            for cell in ws[1]:cell.font=Font(bold=True)
            for item in items:
                ws.append([(v.value if isinstance(v,Enum) else v) for v in (getattr(item,h) for h in headers)])
            ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width=min(max(len(str(c.value or "")) for c in column)+2,45)
        wb.save(target);return target.resolve()
